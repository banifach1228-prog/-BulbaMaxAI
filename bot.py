import os
import json
import time
import base64
import ast
import operator as op
import io
import csv
import re
import threading
from concurrent.futures import ThreadPoolExecutor
import tempfile
import secrets
from difflib import SequenceMatcher
from pathlib import Path

from licenses import (
    activate_license,
    admin_ids,
    block_user,
    create_license,
    get_license_status,
    is_admin,
    release_request,
    reserve_request,
    revoke_user,
    unblock_user,
    user_has_access,
)

import requests

import media_service

BOT_VERSION = "V18.10"
BOT_TOKEN = os.getenv("BOT_TOKEN", "").strip()
API_KEY = os.getenv("API_KEY", "").strip()

BASE_URL = "https://api.baza-ai.org/v1"
TG_URL = f"https://api.telegram.org/bot{BOT_TOKEN}"
DATA_FILE = "v17_memory.json"
LEGACY_DATA_FILES = ("v16_memory.json", "v15_memory.json", "v14_memory.json")

MAX_HISTORY = 24
MAX_CHATS = 30
MAX_REPLY = 3900
MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_IMAGE_BYTES = 10 * 1024 * 1024
MAX_FILE_TEXT = 40000
POLL_TIMEOUT = 25
API_TIMEOUT = 90
RATE_LIMIT_COUNT = 6
RATE_LIMIT_WINDOW = 10

STYLE_PROMPTS = {
    "normal": "Отвечай понятно, точно и по делу. Если есть неопределённость, прямо обозначай её.",
    "short": "Отвечай кратко и по существу. Не добавляй лишнюю воду.",
    "detailed": "Отвечай подробно и структурированно. Объясняй важные шаги и причины, но не растягивай ответ без необходимости.",
}

if not BOT_TOKEN:
    raise RuntimeError("BOT_TOKEN is not set")
if not API_KEY:
    raise RuntimeError("API_KEY is not set")

tg_session = requests.Session()
api_session = requests.Session()
api_session.headers.update({
    "Authorization": f"Bearer {API_KEY}",
    "Content-Type": "application/json",
})

db = {"users": {}, "total_requests": 0, "total_errors": 0, "rules": [], "version": BOT_VERSION}
DB_LOCK = threading.RLock()
AI_INFLIGHT_LOCK = threading.RLock()
AI_INFLIGHT_USERS = set()
TASK_CANCEL_EVENTS = {}
AI_WORKERS = max(2, min(8, int(os.getenv("AI_WORKERS", "4"))))
AI_EXECUTOR = ThreadPoolExecutor(max_workers=AI_WORKERS, thread_name_prefix="bulba-ai")
MEDIA_RESULT_WORKERS = max(1, min(4, int(os.getenv("MEDIA_RESULT_WORKERS", "2"))))
MEDIA_RESULT_EXECUTOR = ThreadPoolExecutor(max_workers=MEDIA_RESULT_WORKERS, thread_name_prefix="bulba-media-result")


def load_db():
    global db
    with DB_LOCK:
        source = DATA_FILE
        if not Path(source).exists():
            for legacy in LEGACY_DATA_FILES:
                if Path(legacy).exists():
                    source = legacy
                    break
        try:
            with open(source, "r", encoding="utf-8") as f:
                data = json.load(f)
            if isinstance(data, dict):
                db.update(data)
        except (OSError, json.JSONDecodeError):
            pass
        db.setdefault("users", {})
        db.setdefault("total_requests", 0)
        db.setdefault("total_errors", 0)
        db.setdefault("rules", [])
        db.setdefault("version", BOT_VERSION)
        db["version"] = BOT_VERSION
        for u in db["users"].values():
            if isinstance(u, dict):
                u.pop("_user_id", None)
                u.pop("_telegram_user", None)


def save_db():
    with DB_LOCK:
        tmp = DATA_FILE + ".tmp"
        try:
            clean = json.loads(json.dumps(db, ensure_ascii=False))
            for u in clean.get("users", {}).values():
                if isinstance(u, dict):
                    u.pop("_user_id", None)
                    u.pop("_telegram_user", None)
            with open(tmp, "w", encoding="utf-8") as f:
                json.dump(clean, f, ensure_ascii=False, indent=2)
            os.replace(tmp, DATA_FILE)
        except OSError as e:
            print("DB save error:", repr(e))


def get_rule_response(text):
    """Backward-compatible alias for the single global-rule engine."""
    return match_global_rule(text)


def add_rule(kind, pattern, response, priority=0, threshold=0.86):
    return add_global_rule(pattern, response, mode=kind, priority=priority, threshold=threshold)


def remove_rule(rule_id):
    return delete_global_rule(rule_id)


def set_rule_enabled(rule_id, enabled):
    try:
        update_global_rule(rule_id, enabled=enabled)
        return True
    except KeyError:
        return False

def default_chat():
    return {
        "history": [],
        "requests": 0,
        "created": int(time.time()),
        "last_prompt": None,
        "last_request": None,
    }


def default_user():
    return {
        "model": "auto",
        "style": "normal",
        "memory": {},
        "chats": {"main": default_chat()},
        "active_chat": "main",
        "requests": 0,
        "errors": 0,
        "rate": [],
        "pinned_chats": [],
        "favorites": [],
        "pending_action": None,
        "notifications": True,
        "media_models": {},
        "agent_mode": False,
    }


def get_user(uid):
    uid = str(uid)
    if uid not in db["users"]:
        db["users"][uid] = default_user()
    u = db["users"][uid]
    # Runtime identity is never persisted inside the user record.
    u.pop("_user_id", None)
    u.pop("_telegram_user", None)
    u.setdefault("model", "auto")
    u.setdefault("style", "normal")
    u.setdefault("memory", {})
    u.setdefault("chats", {"main": default_chat()})
    u.setdefault("active_chat", "main")
    u.setdefault("requests", 0)
    u.setdefault("errors", 0)
    u.setdefault("rate", [])
    u.setdefault("pinned_chats", [])
    u.setdefault("favorites", [])
    u.setdefault("pending_action", None)
    u.setdefault("media_models", {})
    u.setdefault("agent_mode", False)
    if not isinstance(u.get("media_models"), dict):
        u["media_models"] = {}
    u.setdefault("notifications", True)
    if not u["chats"]:
        u["chats"]["main"] = default_chat()
        u["active_chat"] = "main"
    for c in u["chats"].values():
        c.setdefault("history", [])
        c.setdefault("requests", 0)
        c.setdefault("created", int(time.time()))
        c.setdefault("last_prompt", None)
        c.setdefault("last_request", None)
    return u


def get_chat(u):
    name = u.get("active_chat", "main")
    if name not in u["chats"]:
        u["chats"][name] = default_chat()
    return u["chats"][name]


def allowed_request(u):
    now = time.time()
    u["rate"] = [x for x in u.get("rate", []) if now - x < RATE_LIMIT_WINDOW]
    if len(u["rate"]) >= RATE_LIMIT_COUNT:
        return False
    u["rate"].append(now)
    return True


def tg(method, data=None, timeout=40, files=None):
    try:
        r = tg_session.post(
            f"{TG_URL}/{method}",
            data=data or {},
            files=files,
            timeout=timeout,
        )
        if not r.ok:
            print("Telegram HTTP:", r.status_code, r.text[:500])
            return None
        p = r.json()
        if not p.get("ok"):
            print("Telegram API:", p)
            return None
        return p.get("result")
    except Exception as e:
        print("Telegram error:", repr(e))
        return None


def clean_ai_text(text):
    """Normalize common Markdown artifacts because outgoing bot messages are plain text."""
    s = str(text or "...").replace("\r\n", "\n").replace("\r", "\n")
    # Remove fenced-code wrappers but preserve their contents.
    s = re.sub(r"^\s*```[A-Za-z0-9_+-]*\s*\n?", "", s)
    s = re.sub(r"\n?\s*```\s*$", "", s)
    # Headings/bold/strike markers are noise when parse_mode is not used.
    s = re.sub(r"^\s{0,3}#{1,6}\s+", "", s, flags=re.M)
    s = s.replace("**", "").replace("__", "").replace("~~", "")
    # Normalize excessive blank lines for small iPhone screens.
    s = re.sub(r"\n{4,}", "\n\n", s)
    return s.strip() or "..."


def navigation_alias(text):
    t = str(text or "").strip().casefold()
    return {
        "чат": "🤖 Чат", "agent": "🤖 Агент", "агент": "🤖 Агент",
        "модели": "🧠 Модели", "модель": "🧠 Модели", "чаты": "💬 Чаты",
        "генерация": "🎨 Генерация", "память": "💾 Память", "избранное": "⭐ Избранное",
        "статистика": "📊 Статистика", "настройки": "⚙️ Настройки", "лицензия": "🔑 Лицензия",
        "отмена": "❌ Отмена", "cancel": "❌ Отмена",
    }.get(t)


def send_message(chat_id, text, keyboard=None):
    text = clean_ai_text(text)
    chunks = [text[i:i + MAX_REPLY] for i in range(0, len(text), MAX_REPLY)] or ["..."]
    first = None
    for i, chunk in enumerate(chunks):
        data = {"chat_id": chat_id, "text": chunk}
        if keyboard and i == len(chunks) - 1:
            data["reply_markup"] = json.dumps(keyboard, ensure_ascii=False)
        result = tg("sendMessage", data)
        if first is None:
            first = result
    return first


def edit_message(chat_id, message_id, text, keyboard=None):
    if len(text) <= MAX_REPLY:
        data = {"chat_id": chat_id, "message_id": message_id, "text": text}
        if keyboard:
            data["reply_markup"] = json.dumps(keyboard, ensure_ascii=False)
        return tg("editMessageText", data)
    first = text[:MAX_REPLY]
    rest = text[MAX_REPLY:]
    tg("editMessageText", {
        "chat_id": chat_id, "message_id": message_id, "text": first
    })
    return send_message(chat_id, rest, keyboard)


def answer_callback(callback_id, text=None):
    data = {"callback_query_id": callback_id}
    if text:
        data["text"] = text
    tg("answerCallbackQuery", data)


def main_keyboard(user_id=None):
    rows = [
        [{"text": "🤖 Чат"}, {"text": "🤖 Агент"}],
        [{"text": "🧠 Модели"}, {"text": "💬 Чаты"}],
        [{"text": "🎨 Генерация"}],
        [{"text": "💾 Память"}, {"text": "⭐ Избранное"}],
        [{"text": "📊 Статистика"}, {"text": "⚙️ Настройки"}],
    ]
    if user_id is not None and is_admin(user_id):
        rows.append([{"text": "👑 Админ"}, {"text": "🔑 Лицензия"}])
    else:
        rows.append([{"text": "🔑 Лицензия"}])
    return {"keyboard": rows, "resize_keyboard": True, "is_persistent": True}



def answer_keyboard():
    return {"inline_keyboard": [[{"text": "⭐ Сохранить", "callback_data": "favorite:last"}], [{"text": "🔄 Повторить", "callback_data": "retry"}]]}

def retry_keyboard():
    return {"inline_keyboard": [[{"text": "🔄 Повторить", "callback_data": "retry"}]]}


def settings_keyboard(style):
    def mark(v, label):
        return ("✅ " if style == v else "") + label
    return {"inline_keyboard": [
        [{"text": mark("normal", "🧠 Обычно"), "callback_data": "style:normal"}],
        [{"text": mark("short", "⚡ Кратко"), "callback_data": "style:short"}],
        [{"text": mark("detailed", "📚 Подробно"), "callback_data": "style:detailed"}],
        [{"text": "⬅️ Назад", "callback_data": "back"}],
    ]}


def memory_keyboard():
    return {"inline_keyboard": [
        [{"text": "🗑 Очистить память", "callback_data": "memory_clear"}],
        [{"text": "⬅️ Назад", "callback_data": "back"}],
    ]}


def get_models(force=False):
    cache = get_models.cache
    now = time.time()
    if not force and cache["models"] and now - cache["time"] < 60:
        return cache["models"]
    try:
        r = api_session.get(f"{BASE_URL}/models", timeout=20)
        if not r.ok:
            return cache["models"]
        data = r.json()
        models = [x for x in data.get("data", []) if isinstance(x, dict) and x.get("id")]
        models.sort(key=lambda x: x["id"].lower())
        cache["models"], cache["time"] = models, now
        return models
    except Exception as e:
        print("Models error:", repr(e))
        return cache["models"]


get_models.cache = {"models": [], "time": 0}


def is_bad_model(mid):
    s = mid.lower()
    return any(x in s for x in (
        "embedding", "moderation", "tts", "transcrib", "realtime",
        "image-generation", "rerank"
    ))


def rank_models(models, vision=False):
    def score(item):
        s = item["id"].lower()
        if is_bad_model(s):
            return -10000
        n = 0
        if vision and any(x in s for x in ("vision", "vl", "omni", "4o", "multimodal", "gemini", "mistral")):
            n += 100
        if any(x in s for x in ("gpt", "claude", "gemini", "mistral", "deepseek", "qwen", "kimi", "minimax", "glm", "grok")):
            n += 20
        return n
    return sorted(models, key=score, reverse=True)


def model_keyboard(u):
    rows = [[{"text": "🤖 Авто", "callback_data": "model:auto"}]]
    visible_models = [m for m in get_models() if not is_bad_model(m["id"])]
    for i, m in enumerate(visible_models):
        mid = m["id"]
        title = ("✅ " if u["model"] == mid else "") + mid
        rows.append([{"text": title[:55], "callback_data": f"model:{i}"}])
    rows.append([{"text": "⬅️ Назад", "callback_data": "back"}])
    return {"inline_keyboard": rows}


def extract_error(r):
    try:
        d = r.json()
        e = d.get("error", d)
        return str(e.get("message", e) if isinstance(e, dict) else e)
    except Exception:
        return r.text[:800]


def call_ai(messages, model):
    """Exactly one provider call. Model fallback is handled by ai_chat, not here."""
    payload = {"model": model, "messages": messages}
    try:
        r = api_session.post(f"{BASE_URL}/chat/completions", json=payload, timeout=API_TIMEOUT)
        if not r.ok:
            return None, extract_error(r), r.status_code
        d = r.json()
        choices = d.get("choices") or []
        if not choices:
            return None, "API не вернуло choices.", r.status_code
        content = choices[0].get("message", {}).get("content", "")
        if isinstance(content, list):
            content = "\n".join(p.get("text", "") for p in content if isinstance(p, dict) and p.get("type") == "text")
        return str(content or "").strip(), None, r.status_code
    except requests.RequestException as e:
        return None, str(e), 0
    except Exception as e:
        return None, repr(e), 0


def choose_model(u, vision=False, preferred=None):
    models = [m for m in get_models() if not is_bad_model(m["id"])]
    available = {m["id"] for m in models}
    if preferred and preferred in available:
        primary = preferred
    else:
        ranked = rank_models(models, vision=vision)
        primary = ranked[0]["id"] if ranked else None
    if not primary:
        return []
    # Exactly one primary model; one fallback is allowed only after a real API failure.
    ranked = [m["id"] for m in rank_models(models, vision=vision) if m["id"] != primary]
    return [primary] + ranked[:1]


# ---------------- SAFE CALCULATOR ----------------

_BIN = {
    ast.Add: op.add, ast.Sub: op.sub, ast.Mult: op.mul, ast.Div: op.truediv,
    ast.FloorDiv: op.floordiv, ast.Mod: op.mod, ast.Pow: op.pow,
}
_UN = {ast.UAdd: op.pos, ast.USub: op.neg}


def _normalize_rule_text(value):
    return " ".join(str(value or "").casefold().split())[:500]

def get_global_rules(enabled_only=True):
    rules = db.get("rules", [])
    return [r for r in rules if isinstance(r, dict) and (not enabled_only or r.get("enabled", True))]

def match_global_rule(text):
    normalized = _normalize_rule_text(text)
    for rule in sorted(get_global_rules(), key=lambda r: int(r.get("priority", 0)), reverse=True):
        pattern = _normalize_rule_text(rule.get("pattern"))
        if not pattern:
            continue
        mode = rule.get("mode", "contains")
        if mode == "exact":
            matched = normalized == pattern
        elif mode == "similar":
            matched = SequenceMatcher(None, normalized, pattern).ratio() >= float(rule.get("threshold", 0.86))
        else:
            matched = pattern in normalized
        if matched:
            return str(rule.get("response") or "")[:4000]
    return None

def add_global_rule(pattern, response, mode="contains", priority=0, threshold=0.86):
    if mode not in ("exact", "contains", "similar"):
        raise ValueError("mode must be exact, contains or similar")
    rule = {"id": int(time.time() * 1000), "pattern": str(pattern).strip()[:500], "response": str(response).strip()[:4000], "mode": mode, "priority": int(priority), "threshold": max(0.5, min(1.0, float(threshold))), "enabled": True, "created": int(time.time())}
    if not rule["pattern"] or not rule["response"]:
        raise ValueError("pattern and response are required")
    db.setdefault("rules", []).append(rule)
    save_db()
    return rule

def update_global_rule(rule_id, **changes):
    for rule in db.get("rules", []):
        if str(rule.get("id")) == str(rule_id):
            for k in ("pattern", "response", "mode", "priority", "enabled", "threshold"):
                if k in changes and changes[k] is not None:
                    rule[k] = changes[k]
            if rule.get("mode") not in ("exact", "contains", "similar"):
                raise ValueError("Некорректный режим правила.")
            rule["pattern"] = str(rule.get("pattern") or "").strip()[:500]
            rule["response"] = str(rule.get("response") or "").strip()[:4000]
            rule["threshold"] = max(0.5, min(1.0, float(rule.get("threshold", 0.86))))
            save_db()
            return rule
    raise KeyError("Rule not found")

def delete_global_rule(rule_id):
    before = len(db.get("rules", []))
    db["rules"] = [r for r in db.get("rules", []) if str(r.get("id")) != str(rule_id)]
    if len(db["rules"]) == before:
        return False
    save_db()
    return True


def safe_eval(expr):
    expr = expr.replace(",", ".").replace("×", "*").replace("÷", "/")
    tree = ast.parse(expr, mode="eval")

    def ev(node):
        if isinstance(node, ast.Expression):
            return ev(node.body)
        if isinstance(node, ast.Constant) and isinstance(node.value, (int, float)):
            if abs(float(node.value)) > 1e100:
                raise ValueError("Слишком большое число.")
            return node.value
        if isinstance(node, ast.UnaryOp) and type(node.op) in _UN:
            return _UN[type(node.op)](ev(node.operand))
        if isinstance(node, ast.BinOp) and type(node.op) in _BIN:
            a, b = ev(node.left), ev(node.right)
            if isinstance(node.op, ast.Pow) and abs(float(b)) > 100:
                raise ValueError("Слишком большая степень.")
            return _BIN[type(node.op)](a, b)
        raise ValueError("Недопустимое выражение.")
    return ev(tree)


def extract_math(text):
    s = text.strip()
    if len(s) > 150 or not re.search(r"\d", s):
        return None
    candidate = re.sub(r"(?i)^(посчитай|вычисли|сколько будет|calculate)\s*", "", s).strip()
    if not re.fullmatch(r"[0-9\s\+\-\*\/\(\)\.,%×÷]+", candidate):
        return None
    try:
        return safe_eval(candidate)
    except Exception:
        return None


# ---------------- FILE TOOLS ----------------

def decode_bytes(data):
    for enc in ("utf-8", "utf-8-sig", "cp1251", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            pass
    return data.decode("utf-8", errors="replace")


def tg_file(file_id):
    info = tg("getFile", {"file_id": file_id})
    if not info or not info.get("file_path"):
        raise RuntimeError("Не удалось получить файл из Telegram.")
    url = f"https://api.telegram.org/file/bot{BOT_TOKEN}/{info['file_path']}"
    r = tg_session.get(url, timeout=60)
    r.raise_for_status()
    if len(r.content) > MAX_FILE_BYTES:
        raise RuntimeError("Файл слишком большой.")
    return r.content


def read_text_file(name, data):
    suffix = Path(name).suffix.lower()

    if suffix in (".txt", ".md", ".json", ".csv", ".log", ".py", ".js", ".html", ".css"):
        text = decode_bytes(data)
        if suffix == ".json":
            try:
                obj = json.loads(text)
                text = json.dumps(obj, ensure_ascii=False, indent=2)
            except Exception:
                pass
        elif suffix == ".csv":
            try:
                rows = list(csv.reader(io.StringIO(text)))
                text = "\n".join(" | ".join(row) for row in rows)
            except Exception:
                pass
        return text[:MAX_FILE_TEXT]

    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise RuntimeError("Для XLSX нужен пакет openpyxl.")
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        parts = []
        try:
            for ws in wb.worksheets:
                parts.append(f"=== Лист: {ws.title} ===")
                for row in ws.iter_rows(values_only=True):
                    vals = ["" if v is None else str(v) for v in row]
                    if any(vals):
                        parts.append(" | ".join(vals))
                        if sum(len(x) for x in parts) > MAX_FILE_TEXT:
                            break
                if sum(len(x) for x in parts) > MAX_FILE_TEXT:
                    break
        finally:
            wb.close()
        return "\n".join(parts)[:MAX_FILE_TEXT]

    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError:
            raise RuntimeError("Для PDF нужен пакет pypdf.")
        reader = PdfReader(io.BytesIO(data))
        parts = []
        for i, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            if text.strip():
                parts.append(f"=== Страница {i + 1} ===\n{text}")
            if sum(len(x) for x in parts) > MAX_FILE_TEXT:
                break
        return "\n\n".join(parts)[:MAX_FILE_TEXT] or "В PDF не удалось извлечь текст."

    if suffix == ".docx":
        try:
            from docx import Document
        except ImportError:
            raise RuntimeError("Для DOCX нужен пакет python-docx.")
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            parts.append("=== Таблица ===")
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts)[:MAX_FILE_TEXT] or "В DOCX нет читаемого текста."

    raise RuntimeError(
        "Формат пока не поддерживается. Поддерживаются TXT, MD, JSON, CSV, LOG, "
        "PY, JS, HTML, CSS, XLSX, PDF и DOCX."
    )


def image_content(data, mime, question):
    b64 = base64.b64encode(data).decode("ascii")
    return [
        {"type": "text", "text": question},
        {"type": "image_url", "image_url": {"url": f"data:{mime};base64,{b64}"}},
    ]


# ---------------- REAL FILE CREATION ----------------

def make_xlsx(rows, filename):
    try:
        from openpyxl import Workbook
    except ImportError:
        raise RuntimeError("Не установлен openpyxl.")
    wb = Workbook()
    ws = wb.active
    ws.title = "BulbaMaxAI"
    for r in rows:
        ws.append([str(x) if x is not None else "" for x in r])
    for col in ws.columns:
        width = min(max(len(str(cell.value or "")) for cell in col) + 2, 50)
        ws.column_dimensions[col[0].column_letter].width = width
    path = Path(filename)
    wb.save(path)
    return str(path)


def make_csv(rows, filename):
    path = Path(filename)
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        csv.writer(f).writerows(rows)
    return str(path)


def make_chart(rows, filename):
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except ImportError:
        raise RuntimeError("Не установлен matplotlib.")
    if len(rows) < 2:
        raise RuntimeError("Для графика нужно минимум 2 строки данных.")
    labels, values = [], []
    for row in rows[1:]:
        if len(row) < 2:
            continue
        try:
            labels.append(str(row[0]))
            values.append(float(str(row[1]).replace(",", ".")))
        except ValueError:
            continue
    if not values:
        raise RuntimeError("Не удалось найти числовые данные для графика.")
    plt.figure(figsize=(9, 5))
    plt.bar(labels, values)
    plt.xticks(rotation=35, ha="right")
    plt.tight_layout()
    path = Path(filename)
    plt.savefig(path, dpi=160)
    plt.close()
    return str(path)


def make_pdf(text, filename, title="BulbaMaxAI"):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_LEFT
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from xml.sax.saxutils import escape

    doc = SimpleDocTemplate(
        str(filename),
        pagesize=A4,
        rightMargin=40,
        leftMargin=40,
        topMargin=40,
        bottomMargin=40,
    )
    # Register a Unicode font when available so Cyrillic text does not turn into empty boxes.
    font_name = "Helvetica"
    for font_path in (
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/dejavu/DejaVuSans.ttf",
    ):
        if Path(font_path).exists():
            try:
                pdfmetrics.registerFont(TTFont("BulbaUnicode", font_path))
                font_name = "BulbaUnicode"
                break
            except Exception:
                pass

    styles = getSampleStyleSheet()
    styles["Title"].fontName = font_name
    body = ParagraphStyle(
        "BulbaBody",
        parent=styles["BodyText"],
        alignment=TA_LEFT,
        leading=15,
        spaceAfter=8,
        fontName=font_name,
    )
    story = [Paragraph(escape(str(title)), styles["Title"]), Spacer(1, 10)]
    for block in str(text or "").split("\n"):
        block = block.strip()
        if block:
            story.append(Paragraph(escape(block), body))
        else:
            story.append(Spacer(1, 6))
    doc.build(story)
    return str(filename)


def make_docx(text, filename, title="BulbaMaxAI"):
    from docx import Document
    doc = Document()
    doc.add_heading(str(title), level=1)
    for block in str(text or "").split("\n"):
        doc.add_paragraph(block)
    doc.save(str(filename))
    return str(filename)


def send_document(chat_id, path, caption=""):
    with open(path, "rb") as f:
        return tg(
            "sendDocument",
            {"chat_id": chat_id, "caption": caption[:1000]},
            files={"document": (Path(path).name, f)},
            timeout=60,
        )


def send_photo(chat_id, path, caption=""):
    with open(path, "rb") as f:
        return tg(
            "sendPhoto",
            {"chat_id": chat_id, "caption": caption[:1000]},
            files={"photo": (Path(path).name, f)},
            timeout=60,
        )


def parse_rows_from_text(text):
    lines = [x.strip() for x in text.splitlines() if x.strip()]
    rows = []
    for line in lines:
        if "|" in line:
            cells = [x.strip() for x in line.strip("|").split("|")]
        elif ";" in line:
            cells = [x.strip() for x in line.split(";")]
        elif "," in line:
            cells = [x.strip() for x in line.split(",")]
        else:
            cells = re.split(r"\s{2,}", line)
        if cells:
            rows.append(cells)
    return rows


def detect_image_edit_intent(text):
    t = str(text or "").casefold()
    edit_words = (
        "измени", "изменить", "замени", "заменить", "убери", "удали",
        "добавь", "добавить", "перекрась", "перекрасить", "сделай фон",
        "поменяй фон", "отретушируй", "отретушировать", "отредактируй",
        "редактируй", "отрисуй", "рестайлинг", "image edit", "edit image",
        "remove object", "replace object", "change background",
    )
    return any(w in t for w in edit_words)


def _catalog_supports_image_input(item):
    blob = json.dumps(item, ensure_ascii=False).casefold()
    return any(k in blob for k in ("image_url", "image_base64", "image-to-image", "image to image", "input_image"))


def choose_image_edit_model(preferred=None):
    models = media_service.models_for_kind("image")
    if preferred:
        exact = next((m for m in models if m.get("id") == preferred), None)
        if exact and _catalog_supports_image_input(exact):
            return exact
    priority = ("flux-kontext", "flux-2", "recraft", "midjourney")
    for pid in priority:
        item = next((m for m in models if str(m.get("id", "")).casefold() == pid), None)
        if item and _catalog_supports_image_input(item):
            return item
    candidates = [m for m in models if _catalog_supports_image_input(m)]
    return candidates[0] if candidates else None


def media_opts_for_image_edit(model, image_data, mime):
    data_url = f"data:{mime};base64,{base64.b64encode(image_data).decode('ascii')}"
    mid = str(model.get("id") or "").casefold()
    opts = {"image_base64": data_url}
    if mid == "flux-2":
        opts.update({"mode": "Image-to-Image", "aspect_ratio": "1:1", "resolution": "1K"})
    elif mid == "flux-kontext":
        opts.update({"aspect_ratio": "1:1", "tier": "pro"})
    return opts


def local_action_plan(text, has_image=False, has_file=False):
    """Deterministic zero-extra-token router. It never calls an AI model."""
    t = str(text or "").strip().lower()
    if not has_image and not has_file and extract_math(text) is not None:
        return {"action": "calculator"}
    if has_image:
        return {"action": "image_edit" if detect_image_edit_intent(text) else "vision"}
    if has_file:
        return {"action": "analyze_file"}
    if any(x in t for x in ("excel", "xlsx", "таблиц", "таблицу", "таблица для скач", "сделай таблицу")):
        return {"action": "create_xlsx"}
    if "csv" in t:
        return {"action": "create_csv"}
    if any(x in t for x in ("график", "диаграмм", "визуализац")):
        return {"action": "create_chart"}
    if "pdf" in t:
        return {"action": "create_pdf"}
    if any(x in t for x in ("docx", "word-документ", "word документ")):
        return {"action": "create_docx"}
    return {"action": "chat"}


def style_system(u):
    memory = u.get("memory", {})
    memory_text = json.dumps(memory, ensure_ascii=False)[:5000]
    return (
        "Ты BulbaMaxAI. " + STYLE_PROMPTS.get(u.get("style", "normal"), STYLE_PROMPTS["normal"]) +
        "\nИспользуй сохранённую память пользователя только когда она относится к запросу.\n"
        f"Память: {memory_text}"
    )


def ai_chat(u, messages, vision=False):
    preferred = None if u.get("model", "auto") == "auto" else u.get("model")
    candidates = choose_model(u, vision=vision, preferred=preferred)
    if not candidates:
        return None, "Нет доступных моделей."
    last_error = "Не удалось получить ответ."
    for index, model in enumerate(candidates):
        content, err, status = call_ai(messages, model)
        if content:
            return content, None
        last_error = err or last_error
        # Fallback only after an actual failure; never probe multiple models on success.
        if index == 0 and len(candidates) > 1 and status in (0, 400, 404, 408, 409, 429, 500, 502, 503, 504):
            continue
        break
    return None, last_error


def build_history(u):
    h = get_chat(u)["history"]
    return h[-MAX_HISTORY:]


def add_history(u, role, content):
    chat = get_chat(u)
    chat["history"].append({"role": role, "content": content})
    chat["history"] = chat["history"][-MAX_HISTORY:]


def detect_media_intent(text):
    """Legacy no-op kept for compatibility with older integrations."""
    return None


def reserve_license_request(user_id):
    """Reserve one licensed request before expensive work starts."""
    if os.getenv("LICENSE_REQUIRED", "0").strip() != "1" or is_admin(user_id):
        return True
    ok, _ = reserve_request(str(user_id))
    return bool(ok)


def release_license_request(user_id):
    """Return a reserved request when processing fails before completion."""
    if os.getenv("LICENSE_REQUIRED", "0").strip() != "1" or is_admin(user_id):
        return True
    ok, _ = release_request(str(user_id))
    return bool(ok)


def consume_license_after_success(user_id):
    """Backward-compatible alias for older integrations."""
    return reserve_license_request(user_id)


def record_success(u, chat):
    u["requests"] = int(u.get("requests") or 0) + 1
    chat["requests"] = int(chat.get("requests") or 0) + 1
    db["total_requests"] = int(db.get("total_requests") or 0) + 1


AGENT_ACTIONS = {
    "chat", "calculator", "vision", "image_edit", "analyze_file",
    "create_pdf", "create_docx", "create_xlsx", "create_csv", "create_chart",
}


def _extract_json_object(text):
    """Best-effort extraction of one JSON object from a model response."""
    raw = str(text or "").strip()
    if raw.startswith("```"):
        raw = re.sub(r"^```(?:json)?\s*", "", raw, flags=re.I)
        raw = re.sub(r"\s*```$", "", raw)
    try:
        obj = json.loads(raw)
        return obj if isinstance(obj, dict) else None
    except Exception:
        pass
    match = re.search(r"\{.*\}", raw, flags=re.S)
    if not match:
        return None
    try:
        obj = json.loads(match.group(0))
        return obj if isinstance(obj, dict) else None
    except Exception:
        return None


def agent_plan(u, text, image=False, file_text=False, file_name=None):
    """Create a small validated plan. The planner cannot execute tools or actions itself."""
    context = []
    if image:
        context.append("В запросе есть изображение.")
    if file_text:
        context.append(f"В запросе есть файл: {file_name or 'файл'}.")
    planner_messages = [
        {"role": "system", "content": (
            "Ты планировщик BulbaMaxAI. Не решай задачу и не придумывай инструменты. "
            "Верни только JSON без markdown в формате: "
            '{"action":"...","steps":["..."],"goal":"..."}. '
            "action должен быть одним из: chat, calculator, vision, image_edit, analyze_file, "
            "create_pdf, create_docx, create_xlsx, create_csv, create_chart. "
            "steps — максимум 5 коротких шагов. Если задача обычная — chat. "
            "Никаких shell, HTTP, файловых путей, удаления данных или внешних действий."
        )},
        {"role": "user", "content": "\n".join(context + [f"Задача пользователя: {text}"])},
    ]
    answer, error = ai_chat(u, planner_messages, vision=False)
    obj = _extract_json_object(answer) if answer else None
    if not obj:
        return {"action": "chat", "steps": ["Понять задачу", "Дать готовый результат"], "goal": text, "planner_error": error}
    action = str(obj.get("action") or "chat").strip().lower()
    if action not in AGENT_ACTIONS:
        action = "chat"
    steps = obj.get("steps")
    if not isinstance(steps, list):
        steps = []
    steps = [str(x).strip()[:300] for x in steps if str(x).strip()][:5]
    if not steps:
        steps = ["Понять задачу", "Выполнить её", "Проверить результат"]
    goal = str(obj.get("goal") or text).strip()[:1000]
    return {"action": action, "steps": steps, "goal": goal}


def start_agent_request(chat_id, user_id, u, text, image=None, file_text=None, file_name=None, media_ref=None):
    """Run agent mode outside polling. One user request reserves one license unit."""
    uid = str(user_id)
    with AI_INFLIGHT_LOCK:
        if uid in AI_INFLIGHT_USERS:
            send_message(chat_id, "⏳ Предыдущая задача ещё выполняется.", main_keyboard(user_id=user_id))
            return False
        AI_INFLIGHT_USERS.add(uid)
        TASK_CANCEL_EVENTS[uid] = threading.Event()

    def worker():
        try:
            handle_agent_request(chat_id, u, text, image=image, file_text=file_text,
                                 file_name=file_name, media_ref=media_ref, user_id=user_id, cancel_event=TASK_CANCEL_EVENTS.get(uid))
        finally:
            with AI_INFLIGHT_LOCK:
                AI_INFLIGHT_USERS.discard(uid)
                TASK_CANCEL_EVENTS.pop(uid, None)

    AI_EXECUTOR.submit(worker)
    return True


def handle_agent_request(chat_id, u, text, image=None, file_text=None, file_name=None, media_ref=None, user_id=None, cancel_event=None):
    if not allowed_request(u):
        send_message(chat_id, "⏳ Слишком много запросов. Подожди несколько секунд.", main_keyboard(user_id=user_id))
        return
    license_user_id = str(user_id if user_id is not None else chat_id)
    reserved = reserve_license_request(license_user_id)
    if not reserved:
        send_message(chat_id, "⛔ Лимит лицензии исчерпан или доступ недоступен.", main_keyboard(user_id=user_id))
        return
    completed = False
    status_msg = send_message(chat_id, "🤖 Агент: планирую задачу…", {"inline_keyboard":[[{"text":"⏹ Отмена","callback_data":"task:cancel"}]]})
    chat = get_chat(u)
    try:
        if cancel_event and cancel_event.is_set():
            raise RuntimeError("TASK_CANCELLED")
        # Deterministic fast path: don't spend an AI call just to plan arithmetic.
        if not image and not file_text and extract_math(text) is not None:
            plan = {"action": "calculator", "steps": ["Распознать выражение", "Посчитать", "Проверить результат"], "goal": text}
        else:
            plan = agent_plan(u, text, image=bool(image), file_text=bool(file_text), file_name=file_name)
        plan_lines = "\n".join(f"{i+1}. {step}" for i, step in enumerate(plan["steps"]))
        if status_msg:
            edit_message(chat_id, status_msg["message_id"], f"🤖 Агент\n\nПлан:\n{plan_lines}\n\n⏳ Выполняю…")

        action = plan["action"]
        if cancel_event and cancel_event.is_set():
            raise RuntimeError("TASK_CANCELLED")
        if action == "calculator" and not image and not file_text:
            calc = extract_math(text)
            if calc is None:
                action = "chat"
            else:
                answer = f"🧮 Ответ: {calc:g}" if isinstance(calc, float) and calc.is_integer() else f"🧮 Ответ: {calc}"
                add_history(u, "user", text); add_history(u, "assistant", answer)
                chat["last_prompt"] = text; chat["last_request"] = {"kind": "text", "text": text}
                record_success(u, chat); save_db(); completed = True
                edit_message(chat_id, status_msg["message_id"], answer, answer_keyboard()) if status_msg else send_message(chat_id, answer, answer_keyboard())
                return

        # Real image editing path: use a media model that explicitly accepts image input.
        if action == "image_edit" and image:
            selected_id = u.get("media_models", {}).get("image")
            selected = choose_image_edit_model(selected_id)
            if not selected:
                raise RuntimeError("Сейчас нет доступной модели, которая поддерживает редактирование изображения.")
            opts = media_opts_for_image_edit(selected, image.get("data", b""), image.get("mime", "image/jpeg"))
            status_edit = send_message(chat_id, f"🖼 Редактирование запущено.\n\nМодель: {selected.get('name') or selected.get('id')}\n⏳ Обрабатываю…", {"inline_keyboard":[[{"text":"⏹ Отмена","callback_data":"task:cancel"}]]})
            job_id, _ = media_service.create_job("image", text, model=selected.get("id"), opts=opts, user_id=None)
            deadline = time.time() + 180
            result = None
            while time.time() < deadline:
                if cancel_event and cancel_event.is_set():
                    raise RuntimeError("TASK_CANCELLED")
                result = media_service.job_status(job_id)
                if result and result.get("status") == "success":
                    break
                if result and result.get("status") in ("failed", "fail"):
                    raise RuntimeError(str(result.get("error") or result.get("failMsg") or "Редактирование не удалось."))
                time.sleep(3)
            if not result or result.get("status") != "success":
                raise RuntimeError("Редактирование заняло слишком много времени. Попробуй ещё раз позже.")
            claimed = media_service.claim_success(job_id)
            if not claimed or not claimed.get("urls"):
                raise RuntimeError("Модель завершила задачу без изображения.")
            raw, content_type = _download_media_result(claimed["urls"][0], timeout=180)
            filename = _media_filename("image", content_type)
            sent = tg("sendPhoto", {"chat_id": chat_id, "caption": "🖼 Готово"}, files={"photo": (filename, raw, content_type)}, timeout=120)
            if not sent:
                sent = tg("sendDocument", {"chat_id": chat_id, "caption": "🖼 Готово"}, files={"document": (filename, raw, content_type)}, timeout=120)
            if not sent:
                raise RuntimeError("Telegram не принял готовое изображение.")
            chat = get_chat(u)
            add_history(u, "user", text)
            add_history(u, "assistant", "Изображение отредактировано.")
            chat["last_prompt"] = text
            chat["last_request"] = {"kind": "image_edit", "text": text}
            record_success(u, chat)
            save_db(); completed = True
            if status_msg:
                edit_message(chat_id, status_msg["message_id"], "🖼 Изображение готово.", answer_keyboard())
            return

        # The execution layer is intentionally limited to the same safe, local capabilities
        # already supported by the normal request handler. No shell/network/file-deletion tools
        # are exposed to the agent planner.
        execution_prompt = (
            "Ты BulbaMaxAI в режиме агента. Выполни задачу пользователя по проверенному плану.\n"
            f"Цель: {plan['goal']}\n"
            f"План:\n{plan_lines}\n\n"
            "Не утверждай, что сделал действие, если фактически его не сделал. "
            "Если данных недостаточно — скажи, что именно нужно. "
            "Проверь итог перед ответом.\n\n"
            f"Задача пользователя: {text}"
        )
        if action == "create_pdf" and "pdf" not in text.lower():
            execution_prompt += "\nЕсли пользователь явно просил PDF, подготовь содержимое для PDF."
        if action == "create_docx" and "docx" not in text.lower() and "word" not in text.lower():
            execution_prompt += "\nЕсли пользователь явно просил Word, подготовь содержимое для Word."

        # Keep the actual payload type intact for vision/files.
        if cancel_event and cancel_event.is_set():
            raise RuntimeError("TASK_CANCELLED")
        if image:
            content = image_content(image.get("data", b""), image.get("mime", "image/jpeg"), execution_prompt)
            messages = [{"role": "system", "content": style_system(u)}] + build_history(u)
            messages.append({"role": "user", "content": content})
            answer, error = ai_chat(u, messages, vision=True)
        else:
            prompt = execution_prompt
            if file_text:
                prompt += f"\n\nФайл: {file_name or 'файл'}\n\nСодержимое:\n{file_text}"
            messages = [{"role": "system", "content": style_system(u)}] + build_history(u)
            messages.append({"role": "user", "content": prompt})
            answer, error = ai_chat(u, messages, vision=False)

        if not answer:
            raise RuntimeError(error or "Агент не смог получить итоговый ответ.")

        # Execute supported local file actions from the validated plan.
        result_text = answer
        if action in ("create_pdf", "create_docx", "create_xlsx", "create_csv", "create_chart"):
            stem = re.sub(r"[^A-Za-zА-Яа-я0-9_-]+", "_", text[:35]).strip("_") or "bulbamaxai_agent"
            workdir = Path(tempfile.mkdtemp(prefix="bulbamaxai_agent_"))
            if action in ("create_xlsx", "create_csv", "create_chart"):
                table_prompt = (
                    "Преобразуй результат в таблицу для файла. Верни ТОЛЬКО строки с ячейками, "
                    "разделёнными символом |, без markdown-таблицы, заголовков и пояснений.\n\n" + answer
                )
                table_answer, table_error = ai_chat(u, [{"role":"system","content":"Ты форматировщик табличных данных."},{"role":"user","content":table_prompt}], vision=False)
                rows = parse_rows_from_text(table_answer or "")
                if len(rows) < 2:
                    raise RuntimeError(table_error or "Агент не смог подготовить минимум две строки для таблицы.")
                clean_rows = [[str(x).strip() for x in row] for row in rows if row]
                if len(clean_rows) < 2:
                    raise RuntimeError("Недостаточно данных для файла.")
                if action == "create_xlsx":
                    path = workdir / f"{stem}.xlsx"; make_xlsx(clean_rows, path); sent = send_document(chat_id, path, "🤖 Агент создал Excel-файл")
                elif action == "create_csv":
                    path = workdir / f"{stem}.csv"; make_csv(clean_rows, path); sent = send_document(chat_id, path, "🤖 Агент создал CSV-файл")
                else:
                    path = workdir / f"{stem}.png"; make_chart(clean_rows, path); sent = send_photo(chat_id, path, "🤖 Агент создал график")
                if not sent:
                    raise RuntimeError("Telegram не принял созданный файл.")
                result_text = f"Создан файл: {path.name}"
            else:
                path = workdir / f"{stem}{'.pdf' if action == 'create_pdf' else '.docx'}"
                if action == "create_pdf":
                    make_pdf(answer, path, "BulbaMaxAI Agent"); caption = "🤖 Агент создал PDF"
                else:
                    make_docx(answer, path, "BulbaMaxAI Agent"); caption = "🤖 Агент создал Word-документ"
                sent = send_document(chat_id, path, caption)
                if not sent:
                    raise RuntimeError("Telegram не принял созданный документ.")
                result_text = f"Создан файл: {path.name}"

        # Agent output is a successful interaction. Keep only the user request and final result.
        add_history(u, "user", text)
        add_history(u, "assistant", result_text)
        chat["last_prompt"] = text
        last_kind = "image" if image else ("file" if file_text else "text")
        chat["last_request"] = {"kind": last_kind, "text": text,
                                 "file_id": (media_ref or {}).get("file_id") if media_ref else None,
                                 "file_name": file_name}
        record_success(u, chat); save_db(); completed = True
        if status_msg:
            edit_message(chat_id, status_msg["message_id"], "🤖 Агент\n\n" + result_text, answer_keyboard())
        else:
            send_message(chat_id, "🤖 Агент\n\n" + result_text, answer_keyboard())
    except Exception as e:
        if str(e) == "TASK_CANCELLED":
            if status_msg:
                edit_message(chat_id, status_msg["message_id"], "⏹ Задача агента отменена.", main_keyboard(user_id=user_id))
            return
        print("Agent error:", repr(e))
        u["errors"] = int(u.get("errors", 0)) + 1
        db["total_errors"] = int(db.get("total_errors", 0)) + 1
        save_db()
        msg = "❌ Агент не смог завершить задачу. Попробуй сформулировать её иначе или отключи режим агента."
        if status_msg:
            edit_message(chat_id, status_msg["message_id"], msg, main_keyboard(user_id=user_id))
        else:
            send_message(chat_id, msg, main_keyboard(user_id=user_id))
    finally:
        if reserved and not completed:
            release_license_request(license_user_id)
        if "workdir" in locals():
            try:
                for item in workdir.iterdir():
                    item.unlink(missing_ok=True)
                workdir.rmdir()
            except Exception:
                pass


def start_ai_request(chat_id, user_id, u, text, image=None, file_text=None, file_name=None, media_ref=None):
    """Run a potentially long AI/file task outside the Telegram polling loop."""
    uid = str(user_id)
    with AI_INFLIGHT_LOCK:
        if uid in AI_INFLIGHT_USERS:
            send_message(chat_id, "⏳ Предыдущий запрос ещё выполняется. Дождись его завершения.", main_keyboard(user_id=user_id))
            return False
        AI_INFLIGHT_USERS.add(uid)
        TASK_CANCEL_EVENTS[uid] = threading.Event()

    def worker():
        try:
            handle_ai_request(
                chat_id, u, text, image=image, file_text=file_text,
                file_name=file_name, media_ref=media_ref, user_id=user_id, cancel_event=TASK_CANCEL_EVENTS.get(uid)
            )
        finally:
            with AI_INFLIGHT_LOCK:
                AI_INFLIGHT_USERS.discard(uid)
                TASK_CANCEL_EVENTS.pop(uid, None)

    AI_EXECUTOR.submit(worker)
    return True


def handle_ai_request(chat_id, u, text, image=None, file_text=None, file_name=None, media_ref=None, user_id=None, cancel_event=None):
    if not allowed_request(u):
        send_message(chat_id, "⏳ Слишком много запросов. Подожди несколько секунд.", main_keyboard())
        return

    license_user_id = str(user_id if user_id is not None else chat_id)
    reserved = reserve_license_request(license_user_id)
    if not reserved:
        send_message(chat_id, "⛔ Лимит лицензии исчерпан или доступ недоступен.", main_keyboard())
        return
    completed = False

    status = f"🧠 {BOT_VERSION}: анализирую задачу…"
    if image:
        status = "📸 Анализирую изображение…"
    elif file_text:
        status = "📎 Читаю и анализирую файл…"
    status_msg = send_message(chat_id, status, {"inline_keyboard":[[{"text":"⏹ Отмена","callback_data":"task:cancel"}]]})

    chat = get_chat(u)

    try:
        if cancel_event and cancel_event.is_set():
            raise RuntimeError("TASK_CANCELLED")
        # Fast deterministic calculator path.
        calc = extract_math(text)
        if calc is not None and not image and not file_text:
            if isinstance(calc, float) and (calc != calc or abs(calc) == float("inf")):
                raise ValueError("Некорректный числовой результат.")
            answer = f"🧮 Ответ: {calc:g}" if isinstance(calc, float) and calc.is_integer() else f"🧮 Ответ: {calc}"
            add_history(u, "user", text)
            add_history(u, "assistant", answer)
            chat["last_prompt"] = text
            chat["last_request"] = {"kind": "text", "text": text}
            record_success(u, chat)
            save_db()
            completed = True
            if status_msg:
                edit_message(chat_id, status_msg["message_id"], answer, answer_keyboard())
            else:
                send_message(chat_id, answer, answer_keyboard())
            return

        plan = local_action_plan(text, has_image=bool(image), has_file=bool(file_text))
        action = plan.get("action", "chat")
        if cancel_event and cancel_event.is_set():
            raise RuntimeError("TASK_CANCELLED")

        if action in ("create_xlsx", "create_csv", "create_chart", "create_pdf", "create_docx"):
            stem = re.sub(r"[^A-Za-zА-Яа-я0-9_-]+", "_", text[:35]).strip("_") or "bulbamaxai"
            title = "BulbaMaxAI"

            if action in ("create_xlsx", "create_csv", "create_chart"):
                rows = parse_rows_from_text(text)
                if not isinstance(rows, list) or len(rows) < 2:
                    rows = parse_rows_from_text(text)

                # Normalize rows so mixed numeric/string values never break file creation.
                clean_rows = []
                for row in rows if isinstance(rows, list) else []:
                    if isinstance(row, (list, tuple)) and row:
                        clean_rows.append([
                            (x if isinstance(x, (int, float)) else str(x))
                            for x in row
                        ])
                rows = clean_rows

                if len(rows) < 2:
                    action = "chat"
                else:
                    ext = {"create_xlsx": ".xlsx", "create_csv": ".csv", "create_chart": ".png"}[action]
                    path = Path(f"{stem}{ext}")

                    if action == "create_xlsx":
                        make_xlsx(rows, path)
                        sent = send_document(chat_id, path, "📊 Готовая таблица Excel")
                        if not sent:
                            raise RuntimeError("Telegram не принял Excel-файл.")
                    elif action == "create_csv":
                        make_csv(rows, path)
                        sent = send_document(chat_id, path, "📄 Готовый CSV-файл")
                        if not sent:
                            raise RuntimeError("Telegram не принял CSV-файл.")
                    else:
                        make_chart(rows, path)
                        sent = send_photo(chat_id, path, "📈 Готовый график")
                        if not sent:
                            raise RuntimeError("Telegram не принял график.")

                    try:
                        path.unlink(missing_ok=True)
                    except Exception:
                        pass
                    add_history(u, "user", text)
                    add_history(u, "assistant", f"Создан файл: {path.name}")
                    chat["last_prompt"] = text
                    chat["last_request"] = {"kind": "text", "text": text}
                    record_success(u, chat)
                    save_db()
                    completed = True
                    return

            if action in ("create_pdf", "create_docx"):
                # Generate the document body with the normal single-model AI path,
                # then render that answer into the requested file format.
                messages = [{"role": "system", "content": style_system(u)}]
                messages += build_history(u)
                messages.append({"role": "user", "content": text})
                answer, error = ai_chat(u, messages, vision=False)
                if not answer:
                    raise RuntimeError(error or "Не удалось подготовить содержимое документа.")
                ext = ".pdf" if action == "create_pdf" else ".docx"
                path = Path(f"{stem}{ext}")
                if action == "create_pdf":
                    make_pdf(answer, path, title)
                    caption = "📕 Готовый PDF"
                else:
                    make_docx(answer, path, title)
                    caption = "📝 Готовый Word-документ"
                sent = send_document(chat_id, path, caption)
                try:
                    path.unlink(missing_ok=True)
                except Exception:
                    pass
                if not sent:
                    raise RuntimeError("Telegram не принял документ.")
                add_history(u, "user", text)
                add_history(u, "assistant", answer)
                chat["last_prompt"] = text
                chat["last_request"] = {"kind": "text", "text": text}
                record_success(u, chat)
                save_db()
                completed = True
                return

        if image:
            mime = image.get("mime", "image/jpeg")
            content = image_content(image["data"], mime, text)
            messages = [{"role": "system", "content": style_system(u)}]
            messages += build_history(u)
            messages.append({"role": "user", "content": content})
            answer, error = ai_chat(u, messages, vision=True)
        else:
            prompt = text
            if file_text:
                prompt = f"Файл: {file_name}\n\nСодержимое:\n{file_text}\n\nЗадача пользователя:\n{text}"
            messages = [{"role": "system", "content": style_system(u)}]
            messages += build_history(u)
            messages.append({"role": "user", "content": prompt})
            answer, error = ai_chat(u, messages, vision=False)

        if not answer:
            u["errors"] += 1
            db["total_errors"] += 1
            save_db()
            msg = f"❌ Не удалось получить ответ.\n\n{error or 'Неизвестная ошибка'}"
            if status_msg:
                edit_message(chat_id, status_msg["message_id"], msg, main_keyboard())
            else:
                send_message(chat_id, msg, main_keyboard())
            return

        add_history(u, "user", text)
        add_history(u, "assistant", answer)
        chat["last_prompt"] = text
        last_kind = "image" if image else ("file" if file_text else "text")
        chat["last_request"] = {
            "kind": last_kind,
            "text": text,
            "file_id": (media_ref or {}).get("file_id") if media_ref else None,
            "file_name": file_name,
        }
        record_success(u, chat)
        save_db()
        completed = True

        if status_msg:
            edit_message(chat_id, status_msg["message_id"], answer, answer_keyboard())
        else:
            send_message(chat_id, answer, answer_keyboard())

    except Exception as e:
        if str(e) == "TASK_CANCELLED":
            if status_msg:
                edit_message(chat_id, status_msg["message_id"], "⏹ Запрос отменён.", main_keyboard(user_id=user_id))
            return
        print("Request error:", repr(e))
        u["errors"] += 1
        db["total_errors"] += 1
        save_db()
        msg = "❌ Ошибка при обработке запроса. Попробуй ещё раз."
        if status_msg:
            edit_message(chat_id, status_msg["message_id"], msg, main_keyboard())
        else:
            send_message(chat_id, msg, main_keyboard())
    finally:
        if reserved and not completed:
            release_license_request(license_user_id)


def show_stats(chat_id, u):
    send_message(chat_id,
        f"📊 Статистика\n\n"
        f"Запросов: {u['requests']}\n"
        f"Ошибок: {u['errors']}\n"
        f"Чатов: {len(u['chats'])}\n"
        f"Текущий чат: {u['active_chat']}\n"
        f"Всего запросов бота: {db['total_requests']}",
        main_keyboard())


def show_memory(chat_id, u):
    mem = u.get("memory", {})
    if not mem:
        text = "💾 Память пуста."
    else:
        text = "💾 Память:\n\n" + "\n".join(f"• {k}: {v}" for k, v in mem.items())
    send_message(chat_id, text, memory_keyboard())


def show_chats(chat_id, u):
    rows=[]
    for name in u["chats"]:
        mark="📌 " if name in u.get("pinned_chats",[]) else ""
        active="✅ " if name==u["active_chat"] else ""
        rows.append([{"text":active+mark+name[:40],"callback_data":"chat:open:"+name[:45]}])
        rows.append([{"text":"✏️","callback_data":"chat:rename:"+name[:40]},{"text":"📌","callback_data":"chat:pin:"+name[:40]},{"text":"🗑","callback_data":"chat:delete:"+name[:40]}])
    rows.append([{"text":"➕ Новый чат","callback_data":"new_chat"}])
    rows.append([{"text":"⬅️ Назад","callback_data":"back"}])
    send_message(chat_id,"💬 Мои чаты:", {"inline_keyboard":rows})


def new_chat(u):
    if len(u["chats"]) >= MAX_CHATS:
        # Remove the oldest inactive chat.
        candidates = [(c.get("created", 0), n) for n, c in u["chats"].items() if n != u["active_chat"]]
        if candidates:
            _, old = sorted(candidates)[0]
            del u["chats"][old]
        else:
            return u["active_chat"]
    base = "Чат"
    i = 1
    while f"{base} {i}" in u["chats"]:
        i += 1
    name = f"{base} {i}"
    u["chats"][name] = default_chat()
    u["active_chat"] = name
    return name



def media_keyboard():
    return {"inline_keyboard": [
        [{"text": "🖼 Изображение", "callback_data": "media:image"}],
        [{"text": "🎬 Видео", "callback_data": "media:video"}],
        [{"text": "🔊 Голос", "callback_data": "media:tts"}],
        [{"text": "🎵 Музыка", "callback_data": "media:music"}],
        [{"text": "🧊 3D", "callback_data": "media:3d"}],
        [{"text": "⬅️ Назад", "callback_data": "back"}],
    ]}


def _media_model_title(item, selected_id=None):
    name = str(item.get("name") or item.get("id") or "Модель")
    mid = str(item.get("id") or "")
    price = item.get("fromRub")
    price_text = f" · от {price:g} ₽" if price is not None and not item.get("priceUnavailable") else ""
    mark = "✅ " if selected_id == mid else ""
    return (mark + name + price_text)[:55]


def media_model_keyboard(u, kind, page=0):
    models = media_service.models_for_kind(kind)
    selected_id = u.get("media_models", {}).get(kind) or "auto"
    per_page = 8
    pages = max(1, (len(models) + per_page - 1) // per_page)
    page = max(0, min(int(page or 0), pages - 1))
    begin = page * per_page
    visible = models[begin:begin + per_page]

    rows = [[{
        "text": "🤖 Авто" + (" · выбрано" if selected_id == "auto" else ""),
        "callback_data": f"mmdl:{kind}:auto:{page}",
    }]]
    for i, item in enumerate(visible):
        absolute = begin + i
        rows.append([{
            "text": _media_model_title(item, selected_id),
            "callback_data": f"mmdl:{kind}:{absolute}:{page}",
        }])

    nav = []
    if page > 0:
        nav.append({"text": "◀️", "callback_data": f"mmdlpage:{kind}:{page - 1}"})
    nav.append({"text": f"{page + 1}/{pages}", "callback_data": "noop"})
    if page + 1 < pages:
        nav.append({"text": "▶️", "callback_data": f"mmdlpage:{kind}:{page + 1}"})
    if nav:
        rows.append(nav)
    rows.append([{
        "text": "🔄 Обновить список",
        "callback_data": f"mmdlrefresh:{kind}",
    }])
    rows.append([{
        "text": "⬅️ Назад",
        "callback_data": "media:menu",
    }])
    return {"inline_keyboard": rows}


def show_media_models(chat_id, u, kind, message_id=None, page=0):
    kind = media_service.normalize_kind(kind)
    if not kind:
        send_message(chat_id, "❌ Неизвестный тип генерации.", main_keyboard())
        return
    models = media_service.models_for_kind(kind)
    labels = {"image":"🖼 Изображение", "video":"🎬 Видео", "tts":"🔊 Голос", "music":"🎵 Музыка", "3d":"🧊 3D"}
    if not models:
        text = f"{labels.get(kind, '🎨 Генерация')}\n\n⚠️ Сейчас доступных моделей нет.\nНажми «Обновить список»."
        kb = {"inline_keyboard": [[{"text": "🔄 Обновить список", "callback_data": f"mmdlrefresh:{kind}"}], [{"text": "⬅️ Назад", "callback_data": "media:menu"}]]}
        if message_id:
            edit_message(chat_id, message_id, text, kb)
        else:
            send_message(chat_id, text, kb)
        return

    selected = u.get("media_models", {}).get(kind) or "auto"
    selected_item = next((m for m in models if m.get("id") == selected), None)
    if selected != "auto" and not selected_item:
        u.setdefault("media_models", {}).pop(kind, None)
        selected = "auto"
        save_db()
    if selected == "auto":
        selected_text = "🤖 Авто (самая подходящая доступная модель)"
    else:
        selected_text = str(selected_item.get("name") or selected)

    pages = max(1, (len(models) + 7) // 8)
    page = max(0, min(int(page or 0), pages - 1))
    text = (
        f"{labels.get(kind, '🎨 Генерация')}\n\n"
        f"🧠 Выбор модели\n"
        f"Текущая: {selected_text}\n"
        f"Доступно моделей: {len(models)}\n\n"
        "Выбери модель ниже:"
    )
    kb = media_model_keyboard(u, kind, page)
    if message_id:
        edit_message(chat_id, message_id, text, kb)
    else:
        send_message(chat_id, text, kb)

def start_media_prompt(chat_id, user_id, u, kind, model=None):
    kind = media_service.normalize_kind(kind)
    user_id = str(user_id)
    if not media_service.configured():
        send_message(chat_id, "⚠️ Media Studio сейчас недоступна: PLUSVIBE_API_KEY не настроен.", main_keyboard(user_id=user_id))
        return
    models = media_service.models_for_kind(kind)
    if not models:
        send_message(chat_id, "⚠️ Для этого типа генерации сейчас нет доступных моделей.", media_keyboard())
        return
    chosen = str(model or u.get("media_models", {}).get(kind) or "auto")
    if chosen != "auto":
        try:
            selected = media_service.choose_model(kind, chosen)
            chosen = selected["id"]
            u.setdefault("media_models", {})[kind] = chosen
        except Exception:
            chosen = "auto"
            u.setdefault("media_models", {}).pop(kind, None)
    u["pending_action"] = {"type": "media_prompt", "kind": kind, "model": chosen}
    save_db()
    labels = {"image":"🖼 Изображение", "video":"🎬 Видео", "tts":"🔊 Голос", "music":"🎵 Музыка", "3d":"🧊 3D"}
    model_text = "Авто" if chosen == "auto" else chosen
    kb = {"inline_keyboard": [
        [{"text": f"🧠 Модель: {model_text}"[:55], "callback_data": f"media:choose:{kind}"}],
        [{"text": "❌ Отмена", "callback_data": "media:menu"}],
    ]}
    send_message(chat_id, f"{labels.get(kind, '🎨 Генерация')}\n\nМодель: {model_text}\n\nНапиши описание того, что нужно создать.", kb)

def admin_keyboard():
    return {"inline_keyboard": [
        [{"text": "👥 Пользователи", "callback_data": "admin:users"}],
        [{"text": "🔑 Лицензии", "callback_data": "admin:licenses"}],
        [{"text": "📜 Правила", "callback_data": "admin:rules"}],
        [{"text": "📊 Статистика", "callback_data": "admin:stats"}],
        [{"text": "🚫 Блокировки", "callback_data": "admin:blocks"}],
        [{"text": "⚙️ Настройки", "callback_data": "admin:settings"}],
        [{"text": "⬅️ Назад", "callback_data": "back"}],
    ]}


def favorite_keyboard(fid):
    return {"inline_keyboard": [[{"text": "🗑 Удалить", "callback_data": f"favdel:{fid}"}]]}


def show_license(chat_id, user_id):
    send_message(chat_id, "🔑 Лицензия\n\n" + get_license_status(user_id), {
        "inline_keyboard": [
            [{"text": "🔑 Активировать код", "callback_data": "license:activate"}],
            [{"text": "🔄 Обновить", "callback_data": "license:status"}],
        ]
    })


def show_settings(chat_id, u):
    kb = {"inline_keyboard": [
        [{"text": ("✅ " if u["style"] == "normal" else "") + "🧠 Обычно", "callback_data": "style:normal"}],
        [{"text": ("✅ " if u["style"] == "short" else "") + "⚡ Кратко", "callback_data": "style:short"}],
        [{"text": ("✅ " if u["style"] == "detailed" else "") + "📚 Подробно", "callback_data": "style:detailed"}],
        [{"text": ("🔔 Уведомления: ВКЛ" if u.get("notifications", True) else "🔕 Уведомления: ВЫКЛ"), "callback_data": "settings:notifications"}],
        [{"text": "🧹 Очистить текущий чат", "callback_data": "settings:clear"}],
        [{"text": "⬅️ Назад", "callback_data": "back"}],
    ]}
    send_message(chat_id, "⚙️ Настройки\n\nВыбери нужный параметр:", kb)


def show_favorites(chat_id, u, user_id=None):
    favs = u.get("favorites", [])
    if not favs:
        send_message(chat_id, "⭐ Избранное\n\nПока ничего не сохранено.", main_keyboard(user_id=user_id))
        return
    for i, item in enumerate(favs[-20:], 1):
        fid = str(item.get("id", i))
        send_message(chat_id, f"⭐ Ответ #{i}\n\n{item.get('text','')}", favorite_keyboard(fid))


def save_favorite(u, text):
    favs = u.setdefault("favorites", [])
    fid = secrets.token_hex(4)
    favs.append({"id": fid, "text": str(text)[:12000], "created": int(time.time())})
    u["favorites"] = favs[-50:]
    save_db()
    return fid




def _download_media_result(url, timeout=180):
    """Download the provider result ourselves so Telegram does not have to fetch a signed URL."""
    r = requests.get(str(url), timeout=timeout, allow_redirects=True)
    r.raise_for_status()
    data = r.content
    if not data:
        raise RuntimeError("Провайдер вернул пустой файл.")
    content_type = (r.headers.get("Content-Type") or "application/octet-stream").split(";", 1)[0].lower()
    return data, content_type


def _media_filename(kind, content_type):
    ext = {
        "image/jpeg": ".jpg", "image/png": ".png", "image/webp": ".webp",
        "video/mp4": ".mp4", "video/webm": ".webm",
        "audio/mpeg": ".mp3", "audio/mp4": ".m4a", "audio/wav": ".wav",
    }.get(content_type)
    if not ext:
        ext = {"image": ".jpg", "video": ".mp4", "tts": ".mp3", "music": ".mp3", "3d": ".glb"}.get(kind, ".bin")
    return f"bulbamaxai_{kind}{ext}"


def run_media_job(chat_id, user_id, job_id, kind):
    user_id = str(user_id)
    u = get_user(user_id)
    deadline = time.time() + {"image":180,"video":720,"tts":240,"music":360,"3d":420}.get(kind,360)
    while time.time() < deadline:
        job = media_service.job_status(job_id)
        if not job:
            send_message(chat_id, "❌ Медиа-задача не найдена.", main_keyboard(user_id=user_id)); return
        if job.get("status") == "success":
            claimed = media_service.claim_success(job_id)
            if not claimed:
                return
            urls = claimed.get("urls") or []
            if not urls:
                media_service._set_job(job_id, finalized=False)
                send_message(chat_id, "❌ Провайдер завершил генерацию, но не вернул файл.", main_keyboard(user_id=user_id)); return
            url = urls[0]
            try:
                raw, content_type = _download_media_result(url, timeout=180 if kind != "video" else 300)
                filename = _media_filename(kind, content_type)
                if kind == "image":
                    sent = tg("sendPhoto", {"chat_id": chat_id, "caption": "🖼 Готово"}, files={"photo": (filename, raw, content_type)}, timeout=120)
                    if not sent:
                        sent = tg("sendDocument", {"chat_id": chat_id, "caption": "🖼 Готово"}, files={"document": (filename, raw, content_type)}, timeout=120)
                elif kind == "video":
                    sent = tg("sendVideo", {"chat_id": chat_id, "caption": "🎬 Готово"}, files={"video": (filename, raw, content_type)}, timeout=300)
                    if not sent:
                        sent = tg("sendDocument", {"chat_id": chat_id, "caption": "🎬 Готово"}, files={"document": (filename, raw, content_type)}, timeout=300)
                elif kind in ("tts", "music"):
                    sent = tg("sendAudio", {"chat_id": chat_id, "caption": "🔊 Готово"}, files={"audio": (filename, raw, content_type)}, timeout=180)
                    if not sent:
                        sent = tg("sendDocument", {"chat_id": chat_id, "caption": "🎵 Готово"}, files={"document": (filename, raw, content_type)}, timeout=180)
                else:
                    sent = tg("sendDocument", {"chat_id": chat_id, "caption": "🧊 3D результат"}, files={"document": (filename, raw, content_type)}, timeout=180)
            except Exception as exc:
                print("Media result download/send error:", repr(exc))
                sent = None
            if not sent:
                # Allow a later status check/retry to finalize the result again.
                media_service._set_job(job_id, finalized=False)
                send_message(chat_id, "❌ Генерация готова, но файл не удалось отправить в Telegram. Попробуй ещё раз.", main_keyboard(user_id=user_id))
                return
            prompt = str(claimed.get("prompt") or "Медиа-задача")
            add_history(u, "user", prompt)
            add_history(u, "assistant", f"Создано медиа: {url}")
            chat = get_chat(u)
            chat["last_prompt"] = prompt
            chat["last_request"] = {"kind": kind, "text": prompt}
            record_success(u, chat)
            save_db()
            send_message(chat_id, "Готово.\n\nМожно создать ещё один результат.", main_keyboard(user_id=user_id))
            return
        if job.get("status") == "failed":
            send_message(chat_id, f"❌ Генерация не удалась.\n\n{job.get('error','Неизвестная ошибка')}", main_keyboard(user_id=user_id)); return
        time.sleep(2.0)
    send_message(chat_id, "⏳ Генерация ещё выполняется. Проверь результат позже или создай новый запрос.", main_keyboard(user_id=user_id))


def start_media_job(chat_id, user_id, u, kind, prompt, model=None):
    try:
        job_id, selected = media_service.create_job(kind, prompt, model=(None if model in (None, "auto") else model), user_id=user_id)
        u["pending_action"] = None
        save_db()
        send_message(chat_id, f"🎨 Генерация запущена.\n\nМодель: {selected['id']}\n⏳ Статус: выполняется…")
        MEDIA_RESULT_EXECUTOR.submit(run_media_job, chat_id, user_id, job_id, kind)
    except PermissionError as e:
        u["pending_action"] = None; save_db(); send_message(chat_id, f"⛔ {e}", main_keyboard(user_id=user_id))
    except Exception as e:
        u["pending_action"] = None; save_db(); print("Media start error:", repr(e)); send_message(chat_id, "❌ Не удалось запустить генерацию. Попробуй ещё раз.", main_keyboard(user_id=user_id))


def process_message(msg):
    if "chat" not in msg:
        return
    chat_id = msg["chat"]["id"]
    user_id = msg.get("from", {}).get("id", chat_id)
    u = get_user(user_id)

    text = (msg.get("text") or "").strip()
    alias = navigation_alias(text)
    if alias:
        text = alias

    # Reply-keyboard navigation always has priority over a pending input flow.
    # Otherwise a button such as "🤖 Чат" can accidentally become the media prompt.
    navigation_texts = {
        "🤖 Чат", "🤖 Агент", "🧠 Модели", "🧠 Модель", "💬 Чаты",
        "🎨 Генерация", "💾 Память", "⭐ Избранное", "📊 Статистика",
        "⚙️ Настройки", "🔑 Лицензия", "👑 Админ", "🧹 Очистить",
        "🆕 Новый чат", "🤖 Авто",
    }
    pending = u.get("pending_action")
    if text in navigation_texts and pending:
        u["pending_action"] = None
        save_db()
        pending = None

    if text and pending and not text.startswith("/"):
        if pending.get("type") == "media_prompt":
            start_media_job(chat_id, user_id, u, pending.get("kind"), text, pending.get("model"))
            return
        if pending.get("type") == "remember":
            u["memory"][str(int(time.time()))] = text[:1000]
            u["pending_action"] = None
            save_db()
            send_message(chat_id, "💾 Запомнил.", main_keyboard(user_id=user_id))
            return
        if pending.get("type") == "activate":
            ok, message = activate_license(user_id, text)
            u["pending_action"] = None
            save_db()
            send_message(chat_id, message, main_keyboard(user_id=user_id))
            return
        if pending.get("type") == "chat_name":
            name = re.sub(r"[\n\r]", " ", text).strip()[:40] or "Чат"
            if name in u["chats"]: name += " 2"
            u["chats"][name] = default_chat(); u["active_chat"] = name; u["pending_action"] = None; save_db()
            send_message(chat_id, f"✅ Создан чат «{name}».\nТеперь ты находишься в нём.", main_keyboard(user_id=user_id))
            return
        if pending.get("type") == "rename_chat":
            old=pending.get("old")
            name=re.sub(r"[\n\r]", " ", text).strip()[:40] or old
            if old in u["chats"] and name not in u["chats"]:
                u["chats"][name]=u["chats"].pop(old)
                if u["active_chat"]==old: u["active_chat"]=name
                if old in u.get("pinned_chats",[]): u["pinned_chats"]=[name if x==old else x for x in u["pinned_chats"]]
                u["pending_action"]=None; save_db(); send_message(chat_id,f"✏️ Чат переименован в «{name}».",main_keyboard(user_id=user_id))
            else:
                u["pending_action"]=None; save_db(); send_message(chat_id,"❌ Такое название уже занято.",main_keyboard(user_id=user_id))
            return
    caption = (msg.get("caption") or "").strip()

    # Global admin rules are enforced server-side and apply to every user.
    if text and not text.startswith("/"):
        rule_answer = get_rule_response(text)
        if rule_answer:
            send_message(chat_id, rule_answer, main_keyboard())
            return

    # Public access/account commands. AI features are checked below.
    if text.startswith("/myid"):
        send_message(chat_id, f"🆔 Твой Telegram ID: {user_id}", main_keyboard())
        return
    if text.startswith("/profile"):
        status = get_license_status(user_id)
        send_message(chat_id, status, main_keyboard())
        return
    if text.startswith("/activate"):
        code = text[len("/activate"):].strip()
        if not code:
            send_message(chat_id, "🔑 Используй: /activate КОД", main_keyboard())
            return
        ok, message = activate_license(user_id, code)
        send_message(chat_id, message, main_keyboard())
        return
    if text.startswith("/admin"):
        if not is_admin(user_id):
            send_message(chat_id, "⛔ Нет доступа.", main_keyboard())
            return
        send_message(chat_id,
            "👑 Админ-команды:\n\n"
            "/give USER_ID DAYS [REQUESTS] — выдать лицензию\n"
            "/newcode DAYS [REQUESTS] — создать код\n"
            "/license USER_ID — статус\n"
            "/revoke USER_ID — отозвать доступ\n"
            "/block USER_ID — заблокировать\n"
            "/unblock USER_ID — разблокировать\n/rules — глобальные правила\n/rule_add TYPE TEXT RESPONSE — добавить правило\n/rule_del ID — удалить правило",
            main_keyboard())
        return
    if text.startswith("/rules"):
        if not is_admin(user_id):
            send_message(chat_id, "⛔ Нет доступа.", main_keyboard())
            return
        rules = db.get("rules", [])
        if not rules:
            send_message(chat_id, "📜 Глобальных правил пока нет.", main_keyboard())
        else:
            lines = ["📜 Глобальные правила:"]
            for r in rules:
                lines.append(f"\n• {r.get('id')} | {r.get('kind')} | {'ON' if r.get('enabled', True) else 'OFF'}\n  {r.get('pattern')} → {r.get('response')}")
            send_message(chat_id, "\n".join(lines)[:MAX_REPLY], main_keyboard())
        return
    if text.startswith("/rule_add"):
        if not is_admin(user_id):
            send_message(chat_id, "⛔ Нет доступа.", main_keyboard())
            return
        parts = text.split(" ", 3)
        if len(parts) < 4 or parts[1] not in {"exact", "contains", "similar"}:
            send_message(chat_id, "Используй: /rule_add exact|contains|similar ТЕКСТ ОТВЕТ", main_keyboard())
            return
        try:
            r = add_rule(parts[1], parts[2], parts[3])
            send_message(chat_id, f"✅ Правило создано: {r['id']}", main_keyboard())
        except Exception as e:
            send_message(chat_id, f"❌ {e}", main_keyboard())
        return
    if text.startswith("/rule_del"):
        if not is_admin(user_id):
            send_message(chat_id, "⛔ Нет доступа.", main_keyboard())
            return
        parts = text.split()
        if len(parts) != 2 or not remove_rule(parts[1]):
            send_message(chat_id, "❌ Правило не найдено.", main_keyboard())
        else:
            send_message(chat_id, "✅ Правило удалено.", main_keyboard())
        return
    if text.startswith("/newcode"):
        if not is_admin(user_id):
            send_message(chat_id, "⛔ Нет доступа.", main_keyboard())
            return
        parts = text.split()
        if len(parts) < 2:
            send_message(chat_id, "Используй: /newcode DAYS [REQUESTS]", main_keyboard())
            return
        try:
            days = int(parts[1])
            requests_limit = int(parts[2]) if len(parts) > 2 else 0
            code = create_license(days, requests_limit)
            send_message(chat_id, f"🔑 Новый код:\n\n{code}\n\nСрок: {days} дн.\nЛимит запросов: {'без лимита' if requests_limit <= 0 else requests_limit}", main_keyboard())
        except Exception:
            send_message(chat_id, "❌ Формат: /newcode DAYS [REQUESTS]", main_keyboard())
        return
    if text.startswith("/give"):
        if not is_admin(user_id):
            send_message(chat_id, "⛔ Нет доступа.", main_keyboard())
            return
        parts = text.split()
        if len(parts) < 3:
            send_message(chat_id, "Используй: /give USER_ID DAYS [REQUESTS]", main_keyboard())
            return
        try:
            target = parts[1]
            days = int(parts[2])
            requests_limit = int(parts[3]) if len(parts) > 3 else 0
            code = create_license(days, requests_limit)
            ok, message = activate_license(target, code)
            send_message(chat_id, f"{message}\n\nКод: {code}", main_keyboard())
        except Exception:
            send_message(chat_id, "❌ Формат: /give USER_ID DAYS [REQUESTS]", main_keyboard())
        return
    if text.startswith("/license"):
        if not is_admin(user_id):
            send_message(chat_id, "⛔ Нет доступа.", main_keyboard())
            return
        parts = text.split()
        if len(parts) != 2:
            send_message(chat_id, "Используй: /license USER_ID", main_keyboard())
            return
        send_message(chat_id, get_license_status(parts[1]), main_keyboard())
        return
    if text.startswith("/revoke") or text.startswith("/block") or text.startswith("/unblock"):
        if not is_admin(user_id):
            send_message(chat_id, "⛔ Нет доступа.", main_keyboard())
            return
        parts = text.split()
        if len(parts) != 2:
            send_message(chat_id, "Укажи USER_ID.", main_keyboard())
            return
        target = parts[1]
        if text.startswith("/revoke"):
            msg_text = revoke_user(target)
        elif text.startswith("/block"):
            msg_text = block_user(target)
        else:
            msg_text = unblock_user(target)
        send_message(chat_id, msg_text, main_keyboard())
        return

    if text.startswith("/rule"):
        if not is_admin(user_id):
            send_message(chat_id, "⛔ Нет доступа.", main_keyboard()); return
        parts = text.split(" ", 3)
        if len(parts) < 4:
            send_message(chat_id, "Используй: /rule contains|exact ТЕКСТ ОТВЕТ", main_keyboard()); return
        mode = parts[1].strip().lower()
        try:
            add_global_rule(parts[2], parts[3], mode=mode)
            send_message(chat_id, "✅ Глобальное правило добавлено.", main_keyboard())
        except Exception as e:
            send_message(chat_id, f"❌ {e}", main_keyboard())
        return
    if text == "/rules":
        if not is_admin(user_id):
            send_message(chat_id, "⛔ Нет доступа.", main_keyboard()); return
        rules = get_global_rules(enabled_only=False)
        if not rules:
            send_message(chat_id, "📜 Глобальных правил нет.", main_keyboard())
        else:
            send_message(chat_id, "📜 Правила:\n\n" + "\n".join([f"• {r.get('id')}: {r.get('pattern')} → {r.get('response')[:120]}" for r in rules]), main_keyboard())
        return

    if text and not msg.get("photo") and not msg.get("document"):
        rule_response = match_global_rule(text)
        if rule_response:
            send_message(chat_id, rule_response, main_keyboard())
            return

    if os.getenv("LICENSE_REQUIRED", "0").strip() == "1" and not is_admin(user_id):
        allowed, reason = user_has_access(user_id)
        if not allowed:
            send_message(chat_id, reason + "\n\n🔑 Активируй доступ командой /activate КОД.", main_keyboard())
            return

    if text.startswith("/start"):
        send_message(chat_id, f"🤖 BulbaMaxAI {BOT_VERSION} готов.\nПросто отправь задачу — я сам выберу способ выполнения.", main_keyboard())
        return
    if text.startswith("/help"):
        send_message(chat_id, "Просто отправляй текст, фото или поддерживаемый файл. Агент сам определит задачу.\n\nКоманды: /new /clear /stats /memory /remember /forget /models", main_keyboard())
        return
    if text.startswith("/agent"):
        u["agent_mode"] = not bool(u.get("agent_mode", False))
        u["pending_action"] = None
        save_db()
        state = "ВКЛ" if u["agent_mode"] else "ВЫКЛ"
        send_message(chat_id, f"🤖 Режим агента: {state}.\n\n" + ("Отправь задачу — агент составит план и выполнит её." if u["agent_mode"] else "Обычный режим чата включён."), main_keyboard(user_id=user_id))
        return
    if text.startswith("/new"):
        name = new_chat(u)
        save_db()
        send_message(chat_id, f"🆕 Создан чат «{name}».", main_keyboard())
        return
    if text.startswith("/clear"):
        get_chat(u)["history"] = []
        get_chat(u)["last_prompt"] = None
        get_chat(u)["last_request"] = None
        save_db()
        send_message(chat_id, "🧹 Текущий чат очищен.", main_keyboard())
        return
    if text.startswith("/stats"):
        show_stats(chat_id, u)
        return
    if text.startswith("/memory"):
        show_memory(chat_id, u)
        return
    if text.startswith("/remember"):
        arg = text[len("/remember"):].strip()
        if "=" in arg:
            k, v = [x.strip() for x in arg.split("=", 1)]
            if k:
                u["memory"][k] = v[:1000]
                save_db()
                send_message(chat_id, f"💾 Сохранил: {k}", main_keyboard())
            return
        send_message(chat_id, "Используй: /remember ключ=значение", main_keyboard())
        return
    if text.startswith("/forget"):
        k = text[len("/forget"):].strip()
        if k:
            u["memory"].pop(k, None)
            save_db()
        send_message(chat_id, "🗑 Готово.", main_keyboard())
        return
    if text.startswith("/models"):
        send_message(chat_id, "🧠 Доступные модели:", model_keyboard(u))
        return
    if text in ("🤖 Чат", "🤖 Авто"):
        u["model"] = "auto" if text == "🤖 Авто" else u.get("model", "auto")
        u["pending_action"] = None
        save_db()
        send_message(chat_id, "🤖 Режим чата включён.\nОтправь сообщение.", main_keyboard(user_id=user_id))
        return
    if text == "🤖 Агент":
        u["agent_mode"] = not bool(u.get("agent_mode", False))
        u["pending_action"] = None
        save_db()
        state = "ВКЛ" if u["agent_mode"] else "ВЫКЛ"
        send_message(chat_id, f"🤖 Режим агента: {state}.\n\n" + ("Теперь отправляй сложные задачи — агент сначала составит план." if u["agent_mode"] else "Обычный режим чата включён."), main_keyboard(user_id=user_id))
        return
    if text in ("🧠 Модели", "🧠 Модель"):
        send_message(chat_id, "🧠 Выбор модели:", model_keyboard(u))
        return
    if text == "📊 Статистика":
        show_stats(chat_id, u)
        return
    if text == "💾 Память":
        show_memory(chat_id, u)
        return
    if text == "💬 Чаты":
        show_chats(chat_id, u)
        return
    if text == "🆕 Новый чат":
        u["pending_action"] = {"type": "chat_name"}; save_db()
        send_message(chat_id, "💬 Как назвать новый чат?", main_keyboard(user_id=user_id))
        return
    if text == "⭐ Избранное":
        show_favorites(chat_id, u, user_id)
        return
    if text == "🎨 Генерация":
        send_message(chat_id, "🎨 Media Studio\n\nЧто создать?", media_keyboard())
        return
    if text == "🔑 Лицензия":
        show_license(chat_id, user_id)
        return
    if text == "👑 Админ":
        if is_admin(user_id): send_message(chat_id, "👑 Панель администратора", admin_keyboard())
        else: send_message(chat_id, "⛔ Нет доступа.", main_keyboard(user_id=user_id))
        return
    if text == "🧹 Очистить":
        get_chat(u)["history"] = []
        get_chat(u)["last_prompt"] = None
        get_chat(u)["last_request"] = None
        save_db()
        send_message(chat_id, "🧹 Текущий чат очищен.", main_keyboard(user_id=user_id))
        return
    if text == "⚙️ Настройки":
        show_settings(chat_id, u)
        return

    if "photo" in msg:
        photo = msg["photo"][-1]
        try:
            data = tg_file(photo["file_id"])
            if len(data) > MAX_IMAGE_BYTES:
                raise RuntimeError("Изображение слишком большое.")
            request_text = caption or "Что изображено на этом фото? Проанализируй изображение."
            if u.get("agent_mode", False):
                start_agent_request(chat_id, user_id, u, request_text, image={"data": data, "mime": "image/jpeg"}, media_ref={"file_id": photo["file_id"]})
            else:
                start_ai_request(chat_id, user_id, u, request_text, image={"data": data, "mime": "image/jpeg"}, media_ref={"file_id": photo["file_id"]})
        except Exception as e:
            send_message(chat_id, f"❌ Не удалось обработать изображение: {e}", main_keyboard())
        return

    if "document" in msg:
        doc = msg["document"]
        try:
            data = tg_file(doc["file_id"])
            file_text = read_text_file(doc.get("file_name", "file.txt"), data)
            request_text = caption or "Проанализируй этот файл и объясни его содержимое."
            if u.get("agent_mode", False):
                start_agent_request(chat_id, user_id, u, request_text, file_text=file_text, file_name=doc.get("file_name", "file.txt"), media_ref={"file_id": doc["file_id"]})
            else:
                start_ai_request(chat_id, user_id, u, request_text, file_text=file_text, file_name=doc.get("file_name", "file.txt"), media_ref={"file_id": doc["file_id"]})
        except Exception as e:
            send_message(chat_id, f"❌ Не удалось обработать файл: {e}", main_keyboard())
        return

    if text:
        if u.get("agent_mode", False):
            start_agent_request(chat_id, user_id, u, text)
        else:
            start_ai_request(chat_id, user_id, u, text)


def process_callback(q):
    data = q.get("data", "")
    msg = q.get("message") or {}
    chat_id = msg.get("chat", {}).get("id")
    user_id = q.get("from", {}).get("id", chat_id)
    if chat_id is None:
        return
    u = get_user(user_id)

    if data == "back":
        u["pending_action"] = None
        save_db()
        answer_callback(q["id"])
        tg("editMessageText", {"chat_id": chat_id, "message_id": msg.get("message_id"), "text": "🤖 Главное меню"})
        send_message(chat_id, "Готов.", main_keyboard(user_id=user_id))
        return

    if data == "memory_clear":
        u["memory"] = {}
        save_db()
        answer_callback(q["id"], "Память очищена")
        tg("editMessageText", {"chat_id": chat_id, "message_id": msg.get("message_id"), "text": "💾 Память очищена."})
        return

    if data.startswith("style:"):
        value = data.split(":", 1)[1]
        if value not in STYLE_PROMPTS:
            answer_callback(q["id"], "Неизвестный стиль")
            return
        u["style"] = value
        save_db()
        answer_callback(q["id"], "Стиль изменён")
        tg("editMessageText", {"chat_id": chat_id, "message_id": msg.get("message_id"), "text": "⚙️ Стиль ответа изменён."})
        return

    if data.startswith("model:"):
        value = data.split(":", 1)[1]
        if value == "auto":
            u["model"] = "auto"
        else:
            try:
                idx = int(value)
                models = [m for m in get_models() if not is_bad_model(m["id"])]
                if idx < 0 or idx >= len(models):
                    raise ValueError
                u["model"] = models[idx]["id"]
            except Exception:
                answer_callback(q["id"], "Модель уже недоступна")
                return
        save_db()
        answer_callback(q["id"], "Модель выбрана")
        tg("editMessageText", {"chat_id": chat_id, "message_id": msg.get("message_id"), "text": f"🧠 Выбрано: {u['model']}"})
        return

    if data == "new_chat":
        u["pending_action"]={"type":"chat_name"}; save_db(); answer_callback(q["id"], "Введите название"); tg("editMessageText", {"chat_id": chat_id, "message_id": msg.get("message_id"), "text": "💬 Как назвать новый чат?"}); return

    if data.startswith("chat:open:"):
        name = data.split(":", 2)[2]
        if name in u["chats"]:
            u["active_chat"] = name
            save_db()
            answer_callback(q["id"], "Чат выбран")
            edit_message(chat_id, msg.get("message_id"), f"💬 Активен чат «{name}».")
        else:
            answer_callback(q["id"], "Чат не найден")
        return

    if data.startswith("chat:rename:"):
        name = data.split(":", 2)[2]
        if name in u["chats"]:
            u["pending_action"] = {"type": "rename_chat", "old": name}
            save_db()
            answer_callback(q["id"], "Введите новое название")
            send_message(chat_id, "✏️ Отправь новое название чата.", main_keyboard(user_id=user_id))
        else:
            answer_callback(q["id"], "Чат не найден")
        return

    if data.startswith("chat:pin:"):
        name = data.split(":", 2)[2]
        if name not in u["chats"]:
            answer_callback(q["id"], "Чат не найден")
            return
        pins = u.setdefault("pinned_chats", [])
        if name in pins:
            pins.remove(name)
            answer_callback(q["id"], "Откреплён")
        else:
            pins.append(name)
            answer_callback(q["id"], "Закреплён")
        save_db()
        show_chats(chat_id, u)
        return

    if data.startswith("chat:delete:"):
        name = data.split(":", 2)[2]
        if name == "main":
            answer_callback(q["id"], "Главный чат нельзя удалить")
            return
        if name in u["chats"]:
            del u["chats"][name]
            u["pinned_chats"] = [x for x in u.get("pinned_chats", []) if x != name]
            if u.get("active_chat") == name:
                u["active_chat"] = "main"
            save_db()
            answer_callback(q["id"], "Удалён")
            show_chats(chat_id, u)
        else:
            answer_callback(q["id"], "Чат не найден")
        return

    if data == "media:menu":
        u["pending_action"] = None
        save_db()
        answer_callback(q["id"])
        edit_message(chat_id, msg.get("message_id"), "🎨 Генерация", media_keyboard())
        return

    if data == "task:cancel":
        ev = TASK_CANCEL_EVENTS.get(str(user_id))
        if ev:
            ev.set()
            answer_callback(q["id"], "Остановка запрошена")
            edit_message(chat_id, msg.get("message_id"), "⏹ Останавливаю текущую задачу…")
        else:
            answer_callback(q["id"], "Активной задачи нет")
        return

    if data == "noop":
        answer_callback(q["id"])
        return

    if data.startswith("media:choose:"):
        u["pending_action"] = None
        save_db()
        kind = data.split(":", 2)[2]
        answer_callback(q["id"])
        show_media_models(chat_id, u, kind, msg.get("message_id"), 0)
        return

    if data.startswith("media:"):
        kind = data.split(":", 1)[1]
        answer_callback(q["id"], "Выбор модели")
        show_media_models(chat_id, u, kind, msg.get("message_id"), 0)
        return

    if data.startswith("mmdlpage:"):
        u["pending_action"] = None
        save_db()
        _, kind, page = data.split(":", 2)
        answer_callback(q["id"])
        show_media_models(chat_id, u, kind, msg.get("message_id"), int(page))
        return

    if data.startswith("mmdlrefresh:"):
        u["pending_action"] = None
        save_db()
        kind = data.split(":", 1)[1]
        media_service.get_catalog(force=True)
        answer_callback(q["id"], "Список обновлён")
        show_media_models(chat_id, u, kind, msg.get("message_id"), 0)
        return

    if data.startswith("mmdl:"):
        parts = data.split(":")
        if len(parts) not in (3, 4):
            answer_callback(q["id"], "Некорректная кнопка")
            return
        _, kind, value = parts[:3]
        page = int(parts[3]) if len(parts) == 4 and parts[3].isdigit() else 0
        kind = media_service.normalize_kind(kind)
        if not kind:
            answer_callback(q["id"], "Неизвестный тип")
            return
        if value == "auto":
            u.setdefault("media_models", {}).pop(kind, None)
            selected = "auto"
        else:
            try:
                models = media_service.models_for_kind(kind)
                idx = int(value)
                if idx < 0 or idx >= len(models):
                    raise ValueError
                selected = models[idx]["id"]
                media_service.choose_model(kind, selected)
                u.setdefault("media_models", {})[kind] = selected
            except Exception:
                answer_callback(q["id"], "Модель уже недоступна. Обнови список.")
                return
        save_db()
        answer_callback(q["id"], "Модель выбрана")
        start_media_prompt(chat_id, user_id, u, kind, None if selected == "auto" else selected)
        return

    if data == "favorite:last":
        last=get_chat(u).get("history",[])
        answer=next((x.get("content") for x in reversed(last) if x.get("role")=="assistant"),None)
        if answer: save_favorite(u,answer); answer_callback(q["id"],"Сохранено ⭐")
        else: answer_callback(q["id"],"Нечего сохранять")
        return

    if data.startswith("favdel:"):
        fid = data.split(":", 1)[1]
        u["favorites"] = [x for x in u.get("favorites", []) if str(x.get("id")) != fid]
        save_db(); answer_callback(q["id"], "Удалено")
        tg("editMessageText", {"chat_id": chat_id, "message_id": msg.get("message_id"), "text": "🗑 Ответ удалён из избранного."})
        return

    if data == "license:activate":
        u["pending_action"] = {"type": "activate"}; save_db(); answer_callback(q["id"])
        tg("editMessageText", {"chat_id": chat_id, "message_id": msg.get("message_id"), "text": "🔑 Отправь код лицензии следующим сообщением."})
        return
    if data == "license:status":
        answer_callback(q["id"]); edit_message(chat_id, msg.get("message_id"), "🔑 Лицензия\n\n" + get_license_status(user_id), {"inline_keyboard":[[{"text":"🔑 Активировать код","callback_data":"license:activate"}]]})
        return
    if data.startswith("settings:"):
        action = data.split(":",1)[1]
        if action == "notifications":
            u["notifications"] = not u.get("notifications", True); save_db(); answer_callback(q["id"], "Настройка изменена"); show_settings(chat_id,u); return
        if action == "clear":
            get_chat(u)["history"]=[]; get_chat(u)["last_prompt"]=None; get_chat(u)["last_request"]=None; save_db(); answer_callback(q["id"], "Чат очищен"); show_settings(chat_id,u); return

    if data.startswith("chat:delete:"):
        name=data.split(":",2)[2]
        if name != "main" and name in u["chats"]:
            del u["chats"][name]
            if u["active_chat"]==name: u["active_chat"]="main"
            save_db(); answer_callback(q["id"], "Удалён"); show_chats(chat_id,u)
        return

    if data.startswith("admin:"):
        if not is_admin(user_id): answer_callback(q["id"],"Нет доступа"); return
        action=data.split(":",1)[1]; answer_callback(q["id"])
        if action=="stats":
            send_message(chat_id,f"📊 Админ-статистика\n\n👥 Пользователей: {len(db.get('users',{}))}\n💬 Запросов: {db.get('total_requests',0)}\n❌ Ошибок: {db.get('total_errors',0)}\n📜 Правил: {len(db.get('rules',[]))}",admin_keyboard()); return
        if action=="users":
            send_message(chat_id,f"👥 Пользователей: {len(db.get('users',{}))}\n\nДля точечной проверки используй /license USER_ID.",admin_keyboard()); return
        if action=="rules":
            send_message(chat_id,"📜 Управление правилами\n\n/rules\n/rule_add exact|contains|similar ТЕКСТ ОТВЕТ\n/rule_del ID",admin_keyboard()); return
        if action=="licenses":
            send_message(chat_id,"🔑 Лицензии\n\n/newcode DAYS [REQUESTS]\n/give USER_ID DAYS [REQUESTS]\n/license USER_ID\n/revoke USER_ID",admin_keyboard()); return
        if action=="blocks":
            send_message(chat_id,"🚫 Блокировки\n\n/block USER_ID\n/unblock USER_ID",admin_keyboard()); return
        send_message(chat_id,"⚙️ Системные настройки управляются через переменные окружения и конфигурацию проекта.",admin_keyboard()); return

    if data == "retry":
        last = get_chat(u).get("last_request")
        answer_callback(q["id"], "Повторяю")
        if not last:
            send_message(chat_id, "Нет запроса для повтора.", main_keyboard())
            return

        try:
            kind = last.get("kind")
            if kind == "image" and last.get("file_id"):
                data_bytes = tg_file(last["file_id"])
                start_ai_request(
                    chat_id, user_id, u, last.get("text", ""),
                    image={"data": data_bytes, "mime": "image/jpeg"},
                    media_ref={"file_id": last["file_id"]},
                )
            elif kind == "file" and last.get("file_id"):
                data_bytes = tg_file(last["file_id"])
                name = last.get("file_name") or "file.txt"
                file_text = read_text_file(name, data_bytes)
                start_ai_request(
                    chat_id, user_id, u, last.get("text", ""),
                    file_text=file_text,
                    file_name=name,
                    media_ref={"file_id": last["file_id"]},
                )
            else:
                start_ai_request(chat_id, user_id, u, last.get("text", ""))
        except Exception as e:
            send_message(chat_id, f"❌ Не удалось повторить запрос: {e}", main_keyboard())
        return


def configure_telegram_menu():
    # Reset the Telegram menu button to the standard commands button.
    result=tg("setChatMenuButton", {"menu_button": json.dumps({"type":"commands"})})
    if result is None:
        print("Telegram menu button configuration skipped/failed")


def main():
    load_db()
    configure_telegram_menu()
    print(f"BulbaMaxAI {BOT_VERSION} started")
    offset = None
    while True:
        try:
            params = {"timeout": POLL_TIMEOUT}
            if offset is not None:
                params["offset"] = offset
            updates = tg("getUpdates", params, timeout=POLL_TIMEOUT + 10)
            if not updates:
                continue
            for update in updates:
                offset = update["update_id"] + 1
                try:
                    if "callback_query" in update:
                        process_callback(update["callback_query"])
                    elif "message" in update:
                        process_message(update["message"])
                except Exception as e:
                    print("Update error:", repr(e))
        except Exception as e:
            print("Main loop error:", repr(e))
            time.sleep(3)


if __name__ == "__main__":
    main()
